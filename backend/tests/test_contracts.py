import io


def _make_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Service Agreement", level=1)
    doc.add_paragraph(
        "1. Scope of Services\n"
        "The Service Provider agrees to provide consulting services as described in Exhibit A."
    )
    doc.add_paragraph(
        "2. Payment Terms\n"
        "The Client shall pay the Service Provider within 30 days of invoice receipt."
    )
    doc.add_paragraph(
        "3. Termination\n"
        "Either party may terminate this agreement with 30 days written notice."
    )
    doc.add_paragraph(
        "4. Indemnification\n"
        "The Service Provider shall indemnify the Client against all claims arising from negligence."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_contract(auth_client):
    docx_bytes = _make_docx_bytes()
    response = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("test_contract.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["title"] == "test_contract"
    assert data["source"] == "upload"
    assert data["status"] == "reviewed"
    assert data["overall_risk_score"] is not None


def test_upload_unsupported_type(auth_client):
    response = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("test.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 415


def test_list_contracts(auth_client):
    docx_bytes = _make_docx_bytes()
    auth_client.post(
        "/api/contracts/upload",
        files={"file": ("c1.docx", docx_bytes, "application/octet-stream")},
    )
    response = auth_client.get("/api/contracts")
    assert response.status_code == 200
    data = response.json()
    assert data["total"] >= 1
    assert len(data["contracts"]) >= 1


def test_get_contract(auth_client):
    docx_bytes = _make_docx_bytes()
    upload = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("c2.docx", docx_bytes, "application/octet-stream")},
    )
    contract_id = upload.json()["id"]
    response = auth_client.get(f"/api/contracts/{contract_id}")
    assert response.status_code == 200
    assert response.json()["id"] == contract_id


def test_get_clauses(auth_client):
    docx_bytes = _make_docx_bytes()
    upload = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("c3.docx", docx_bytes, "application/octet-stream")},
    )
    contract_id = upload.json()["id"]
    response = auth_client.get(f"/api/contracts/{contract_id}/clauses")
    assert response.status_code == 200
    clauses = response.json()
    assert len(clauses) > 0
    assert "risk_level" in clauses[0]


def test_get_review(auth_client):
    docx_bytes = _make_docx_bytes()
    upload = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("c4.docx", docx_bytes, "application/octet-stream")},
    )
    contract_id = upload.json()["id"]
    response = auth_client.get(f"/api/contracts/{contract_id}/review")
    assert response.status_code == 200
    data = response.json()
    assert "overall_risk_score" in data
    assert "clauses" in data


def test_delete_contract(auth_client):
    docx_bytes = _make_docx_bytes()
    upload = auth_client.post(
        "/api/contracts/upload",
        files={"file": ("del.docx", docx_bytes, "application/octet-stream")},
    )
    contract_id = upload.json()["id"]
    response = auth_client.delete(f"/api/contracts/{contract_id}")
    assert response.status_code == 204

    response = auth_client.get(f"/api/contracts/{contract_id}")
    assert response.status_code == 404


def test_contract_not_found(auth_client):
    response = auth_client.get("/api/contracts/99999")
    assert response.status_code == 404
