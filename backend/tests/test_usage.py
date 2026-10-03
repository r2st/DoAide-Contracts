import io

from tests.conftest import TEST_EMAIL, TEST_PASSWORD


def _make_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("1. Simple clause for testing.")
    doc.add_paragraph("2. Another clause for testing.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_free_tier_review_limit(client, db_session):
    response = client.post(
        "/api/auth/register",
        json={"email": "limit@example.com", "password": "strongpassword", "name": "Limit User"},
    )
    assert response.status_code == 201
    token = response.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = _make_docx_bytes()

    for i in range(3):
        resp = client.post(
            "/api/contracts/upload",
            files={"file": (f"contract_{i}.docx", docx_bytes, "application/octet-stream")},
            headers=headers,
        )
        assert resp.status_code == 201, f"Upload {i} failed: {resp.text}"

    resp = client.post(
        "/api/contracts/upload",
        files={"file": ("contract_4.docx", docx_bytes, "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 402


def test_free_tier_generation_limit(client, db_session):
    response = client.post(
        "/api/auth/register",
        json={"email": "genlimit@example.com", "password": "strongpassword", "name": "Gen User"},
    )
    assert response.status_code == 201
    token = response.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    template_resp = client.post(
        "/api/templates/custom",
        json={
            "name": "Simple Template",
            "category": "nda",
            "template_body": "Agreement between {{party_a}} and {{party_b}}.",
        },
        headers=headers,
    )
    assert template_resp.status_code == 201
    template_id = template_resp.json()["id"]

    for i in range(2):
        resp = client.post(
            "/api/contracts/generate",
            json={
                "template_id": template_id,
                "title": f"Contract {i}",
                "field_values": {"party_a": "A", "party_b": "B"},
            },
            headers=headers,
        )
        assert resp.status_code == 201, f"Generation {i} failed: {resp.text}"

    resp = client.post(
        "/api/contracts/generate",
        json={
            "template_id": template_id,
            "title": "Contract 3",
            "field_values": {"party_a": "A", "party_b": "B"},
        },
        headers=headers,
    )
    assert resp.status_code == 402
